import json
import tarfile
from pathlib import Path

import pytest

from manim_engineer.core import Modality, SegmentKind
from manim_engineer.ingest import (
    Cue,
    arxiv_id,
    detect_kind,
    document_from_latex,
    document_from_markdown,
    ingest,
    is_arxiv,
    load_audio,
    load_document,
    load_pdf_deck,
    load_pdf_paper,
    load_pptx,
    load_transcript,
    load_youtube,
    save_document,
    video_id,
)

MD = """# Softmax temperature

Softmax turns logits into probabilities.

$$
p_i = \\frac{e^{z_i/T}}{\\sum_j e^{z_j/T}}
$$

Lower temperature sharpens the distribution. Higher temperature flattens it.

## Code

```python
import numpy as np
```

Final paragraph.
"""


def test_markdown_segments_headings_equations_code():
    doc = document_from_markdown(MD, source="x.md", title="t")
    kinds = [s.kind for s in doc.segments]
    assert kinds == [
        SegmentKind.HEADING,
        SegmentKind.BODY,
        SegmentKind.EQUATION,
        SegmentKind.BODY,
        SegmentKind.HEADING,
        SegmentKind.CODE,
        SegmentKind.BODY,
    ]
    eq = doc.equations()[0]
    assert eq.latex.startswith("p_i =")
    assert eq.span.para_id == "sec1.eq1"
    assert doc.segments[3].span.para_id == "sec1.p2"


def test_transcript_json_groups_cues_with_spans(tmp_path: Path):
    rows = [{"start": i * 4.0, "duration": 4.0, "text": f"Sentence number {i}."} for i in range(20)]
    p = tmp_path / "t.json"
    p.write_text(json.dumps({"snippets": rows}))
    doc = load_transcript(p)
    assert doc.modality == Modality.TRANSCRIPT
    assert 2 <= len(doc.segments) <= 5
    first = doc.segments[0]
    assert first.span.t_start_s == 0.0 and first.span.t_end_s >= 24.0
    assert "Sentence number 0." in first.text
    save_document(doc, tmp_path / "doc.json")
    assert load_document(tmp_path / "doc.json") == doc


def test_transcript_vtt_and_plain(tmp_path: Path):
    vtt = "WEBVTT\n\n00:00.000 --> 00:03.000\nHello <b>there</b>.\n\n00:03.000 --> 00:06.000\nSecond cue.\n"
    p = tmp_path / "a.vtt"
    p.write_text(vtt)
    doc = load_transcript(p)
    assert "Hello there." in doc.segments[0].text and "<b>" not in doc.segments[0].text
    plain = tmp_path / "b.txt"
    plain.write_text("[00:10] First line here.\n[01:05] Second line here.\n")
    doc2 = load_transcript(plain)
    assert doc2.segments[0].span.t_start_s == 10.0
    assert detect_kind(str(plain)) == "transcript"


def test_pptx_slides_and_notes(tmp_path: Path):
    from pptx import Presentation

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Attention"
    s.placeholders[1].text = "Q K V"
    s.notes_slide.notes_text_frame.text = "Explain scaled dot product here."
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "Second"
    p = tmp_path / "deck.pptx"
    prs.save(str(p))
    doc = load_pptx(p)
    assert doc.modality == Modality.SLIDES
    assert doc.segments[0].span.slide_index == 1 and doc.segments[0].kind == SegmentKind.SLIDE_TEXT
    assert any(s.kind == SegmentKind.SLIDE_NOTES and "scaled dot" in s.text for s in doc.segments)
    assert doc.segments[-1].span.slide_index == 2


def test_pdf_deck_and_paper(tmp_path: Path):
    import pymupdf

    deck = pymupdf.open()
    for i in range(3):
        page = deck.new_page(width=960, height=540)
        page.insert_text((72, 100), f"Slide {i + 1} title", fontsize=28)
        page.insert_text(
            (72, 160), "Some bullet text about attention heads and layers.", fontsize=14
        )
    dp = tmp_path / "deck.pdf"
    deck.save(str(dp))
    doc = ingest(str(dp), workdir=tmp_path / "wd")
    assert doc.modality == Modality.SLIDES
    assert [s.span.slide_index for s in doc.segments] == [1, 2, 3]
    assert doc.assets and Path(next(iter(doc.assets.values()))).exists()
    paper = pymupdf.open()
    for i in range(2):
        page = paper.new_page(width=612, height=792)
        y = 80
        for k in range(12):
            page.insert_text(
                (72, y),
                f"Paragraph {k} on page {i + 1}: " + "lorem ipsum dolor sit amet " * 4,
                fontsize=10,
            )
            y += 50
    pp = tmp_path / "paper.pdf"
    paper.save(str(pp))
    doc2 = ingest(str(pp))
    assert doc2.modality == Modality.MANUSCRIPT
    assert doc2.segments[0].span.para_id.startswith("p1.")
    assert len(doc2.segments) >= 12
    doc3 = load_pdf_paper(pp, max_pages=1)
    assert all(s.span.para_id.startswith("p1.") for s in doc3.segments)
    doc4 = load_pdf_deck(dp, max_pages=2)
    assert len(doc4.segments) == 2


def test_audio_with_injected_transcriber(tmp_path: Path):
    wav = tmp_path / "talk.wav"
    wav.write_bytes(b"RIFF")
    doc = load_audio(
        wav,
        transcriber=lambda p: [
            Cue(0, 3, "Vectors are arrows."),
            Cue(3, 6, "They add tip to tail."),
        ],
    )
    assert doc.modality == Modality.AUDIO
    assert doc.segments[0].span.t_end_s == 6.0
    assert "tip to tail" in doc.segments[0].text


def test_youtube_captions_then_audio_fallback(tmp_path: Path):
    assert video_id("https://www.youtube.com/watch?v=fNk_zzaMoSs&t=1") == "fNk_zzaMoSs"
    assert video_id("https://youtu.be/fNk_zzaMoSs") == "fNk_zzaMoSs"
    with pytest.raises(ValueError):
        video_id("https://example.com")
    doc = load_youtube("https://youtu.be/fNk_zzaMoSs", captions=lambda v: [Cue(0, 2, "hi there.")])
    assert doc.modality == Modality.VIDEO and doc.meta["path"] == "captions"

    def boom(v):
        raise RuntimeError("blocked")

    def fake_audio(v, wd: Path):
        f = wd / f"{v}.mp3"
        f.write_bytes(b"ID3")
        return f

    doc2 = load_youtube(
        "fNk_zzaMoSs",
        captions=boom,
        audio=fake_audio,
        transcriber=lambda p: [Cue(0, 2, "fallback works.")],
        workdir=tmp_path,
    )
    assert doc2.meta["path"] == "audio+asr" and "blocked" in doc2.meta["captions_error"]
    assert doc2.segments[0].text == "fallback works."


TEX = r"""
\documentclass{article}
\newcommand{\sm}{\mathrm{softmax}}
\DeclareMathOperator{\KL}{KL}
\title{Speculative Decoding}
\begin{document}

\section{Method}

The draft model proposes tokens and the target model \emph{verifies} them~\cite{leviathan2023}.

\begin{equation}\label{eq:accept}
\alpha = \min\left(1, \frac{p(x)}{q(x)}\right)
\end{equation}

Acceptance is cheap to evaluate.

\section{Results}

\[
\mathbb{E}[N] = \frac{1 - \alpha^{k+1}}{1 - \alpha}
\]

\end{document}
"""


def test_latex_yields_sections_equation_ledger_and_macros():
    doc = document_from_latex(
        TEX, source="https://arxiv.org/abs/2211.17192", arxiv_id_="2211.17192"
    )
    assert doc.title == "Speculative Decoding"
    assert [s.kind for s in doc.segments] == [
        SegmentKind.HEADING,
        SegmentKind.BODY,
        SegmentKind.EQUATION,
        SegmentKind.BODY,
        SegmentKind.HEADING,
        SegmentKind.EQUATION,
    ]
    eqs = doc.equations()
    # equations are carried verbatim, minus \label, so the plan never re-types them
    assert eqs[0].latex == r"\alpha = \min\left(1, \frac{p(x)}{q(x)}\right)"
    assert eqs[0].span.para_id == "sec1.eq1"
    # equation numbering is document-global while paragraph numbering resets per
    # section, matching document_from_markdown; do not "fix" this to sec2.eq1
    assert eqs[1].span.para_id == "sec2.eq2"
    ledger = doc.meta["equation_ledger"]
    assert [e["label"] for e in ledger] == ["eq:accept", "sec2.eq2"]
    assert ledger[0]["env"] == "equation" and ledger[1]["env"] == "display"
    assert doc.meta["has_equation_ledger"] is True
    assert r"\sm" in doc.meta["macros"] and r"\KL" in doc.meta["macros"]
    # prose is readable: markup expanded, citations dropped
    assert "verifies" in doc.segments[1].text and "leviathan2023" not in doc.segments[1].text


def test_arxiv_id_forms_and_detection(tmp_path: Path, monkeypatch):
    assert arxiv_id("2211.17192") == "2211.17192"
    assert arxiv_id("arXiv:2211.17192v2") == "2211.17192v2"
    assert arxiv_id("https://arxiv.org/abs/2211.17192") == "2211.17192"
    assert arxiv_id("https://arxiv.org/pdf/2211.17192v3") == "2211.17192v3"
    assert is_arxiv("2211.17192") and not is_arxiv("notes.md")
    with pytest.raises(ValueError):
        arxiv_id("https://example.com/paper")
    assert detect_kind("arXiv:2211.17192") == "arxiv"
    # an existing local file must never be mistaken for a paper id and fetched
    monkeypatch.chdir(tmp_path)
    Path("2211.17192").write_text("local")
    with pytest.raises(ValueError):
        detect_kind("2211.17192")


def test_arxiv_ingest_flattens_inputs_with_injected_fetch(tmp_path: Path):
    src = tmp_path / "paper"
    src.mkdir()
    (src / "main.tex").write_text(
        "\\documentclass{article}\n"
        "\\newcommand{\\sm}{\\mathrm{softmax}}\n"
        "\\title{Draft and Verify}\n"
        "\\begin{document}\n"
        "\\section{Method}\n\n"
        "% a comment that must not survive\n"
        "Draft then verify.\n\n"
        "\\input{accept}\n\n"
        "\\end{document}\n"
    )
    (src / "accept.tex").write_text(
        "\\begin{equation}\\label{eq:acc}\n\\alpha = \\min(1, p/q)\n\\end{equation}\n"
    )
    archive = tmp_path / "src.tar.gz"
    with tarfile.open(archive, "w:gz") as tf:
        for f in sorted(src.iterdir()):
            tf.add(f, arcname=f.name)

    doc = ingest("arXiv:2211.17192", workdir=tmp_path / "wd", fetch=lambda aid, wd: archive)
    assert doc.modality == Modality.MANUSCRIPT
    assert doc.title == "Draft and Verify"
    assert doc.meta["arxiv_id"] == "2211.17192"
    assert "comment that must not survive" not in doc.text
    eqs = doc.equations()
    assert len(eqs) == 1 and r"\alpha" in (eqs[0].latex or "")
    assert doc.meta["equation_ledger"][0]["label"] == "eq:acc"
    assert r"\sm" in doc.meta["macros"]
    save_document(doc, tmp_path / "doc.json")
    assert load_document(tmp_path / "doc.json") == doc


PREAMBLE_TEX = r"""
\documentclass{article}
\usepackage{amsmath,mathtools}
\usepackage[dvipsnames]{xcolor}
\usepackage{hyperref}
\newcommand{\loss}[1]{
  \mathcal{L}\left(#1\right)
}
\newcommand{\secret}{\input{/etc/passwd}}
\begin{document}
\section{Loss}

We minimise the loss.

\begin{equation}
\loss{\theta} \coloneqq \sum_i \ell_i
\end{equation}

\end{document}
"""


def test_latex_keeps_whole_definitions_and_allowlisted_packages():
    doc = document_from_latex(PREAMBLE_TEX, source="https://arxiv.org/abs/2603.03251")
    assert "\\newcommand{\\loss}[1]{\n  \\mathcal{L}\\left(#1\\right)\n}" in doc.meta["macros"]
    assert doc.meta["tex_packages"] == ["mathtools", "xcolor"]
    assert doc.meta["tex_packages_ignored"] == ["hyperref"]
    dropped = doc.meta["macros_dropped"]
    assert len(dropped) == 1
    assert "secret" in dropped[0]["source"] and "input" in dropped[0]["reason"]
    assert "secret" not in doc.meta["macros"]


def test_arxiv_ingest_reads_definitions_from_the_papers_own_sty(tmp_path: Path):
    src = tmp_path / "paper"
    src.mkdir()
    (src / "main.tex").write_text(
        "\\documentclass{article}\n"
        "\\usepackage{amsmath,notation}\n"
        "\\begin{document}\n"
        "\\section{Inner products}\n\n"
        "\\begin{equation}\\langle \\vx, \\vy \\rangle\\end{equation}\n\n"
        "\\end{document}\n"
    )
    (src / "notation.sty").write_text(
        "\\ProvidesPackage{notation}\n"
        "\\RequirePackage{bm}\n"
        "% vectors are bold\n"
        "\\newcommand{\\vx}{\\bm{x}}\n"
        "\\newcommand{\\vy}{\\bm{y}}\n"
        "\\def\\@internal{x}\n"
    )
    archive = tmp_path / "src.tar.gz"
    with tarfile.open(archive, "w:gz") as tf:
        for f in sorted(src.iterdir()):
            tf.add(f, arcname=f.name)

    doc = ingest("arXiv:2603.03251", workdir=tmp_path / "wd", fetch=lambda aid, wd: archive)
    assert r"\newcommand{\vx}{\bm{x}}" in doc.meta["macros"]
    assert r"\newcommand{\vy}{\bm{y}}" in doc.meta["macros"]
    assert doc.meta["tex_packages"] == ["bm"]
    assert "@internal" not in doc.meta["macros"]
    assert "vectors are bold" not in doc.meta["macros"]


def test_flatten_inlines_each_local_sty_once(tmp_path: Path):
    from manim_engineer.ingest.arxiv import flatten

    # a vendored package that names itself in comments and requires itself
    (tmp_path / "main.tex").write_text("\\usepackage{notation}\n\\usepackage{notation}\n")
    (tmp_path / "notation.sty").write_text(
        "% load it with \\usepackage{notation} \\usepackage{notation}\n"
        "\\RequirePackage{notation}\n"
        "\\newcommand{\\R}{\\mathbb{R}}\n"
    )
    flat = flatten(tmp_path / "main.tex", tmp_path)
    assert flat.count("\\newcommand{\\R}") == 1
