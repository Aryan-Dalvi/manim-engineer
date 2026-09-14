"""Slide adapters: PPTX (text + speaker notes) and PDF decks (text + page rasters)."""

from __future__ import annotations

import re
from pathlib import Path

from manim_engineer.core import Document, Modality, Segment, SegmentKind, SourceSpan
from manim_engineer.ingest.common import doc_id_for


def load_pptx(path: str | Path, title: str | None = None) -> Document:
    from pptx import Presentation

    path = Path(path)
    doc_id = doc_id_for(str(path))
    prs = Presentation(str(path))
    segments: list[Segment] = []
    for idx, slide in enumerate(prs.slides, start=1):
        texts: list[str] = []
        shape_ids: list[str] = []
        for shape in sorted(slide.shapes, key=lambda s: ((s.top or 0), (s.left or 0))):
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                texts.append(re.sub(r"\s+", " ", shape.text_frame.text).strip())
                shape_ids.append(str(shape.shape_id))
            elif getattr(shape, "has_table", False):
                rows = []
                for r in shape.table.rows:
                    rows.append(" | ".join(c.text.strip() for c in r.cells))
                texts.append("\n".join(rows))
                shape_ids.append(str(shape.shape_id))
        if texts:
            segments.append(
                Segment(
                    seg_id=f"s{len(segments):04d}",
                    text="\n".join(texts),
                    kind=SegmentKind.SLIDE_TEXT,
                    span=SourceSpan(
                        modality=Modality.SLIDES,
                        doc_id=doc_id,
                        slide_index=idx,
                        shape_ids=shape_ids,
                    ),
                )
            )
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                segments.append(
                    Segment(
                        seg_id=f"s{len(segments):04d}",
                        text=notes,
                        kind=SegmentKind.SLIDE_NOTES,
                        span=SourceSpan(modality=Modality.SLIDES, doc_id=doc_id, slide_index=idx),
                    )
                )
    t = (
        title
        or (
            prs.slides[0].shapes.title.text
            if len(prs.slides) and prs.slides[0].shapes.title
            else None
        )
        or path.stem
    )
    return Document(
        doc_id=doc_id,
        title=t,
        modality=Modality.SLIDES,
        source=str(path),
        segments=segments,
        meta={"slides": len(prs.slides)},
    )


def pdf_is_deck(path: str | Path) -> bool:
    import pymupdf

    pdf = pymupdf.open(str(path))
    if pdf.page_count == 0:
        return False
    landscape = sum(1 for p in pdf if p.rect.width > p.rect.height)
    avg_chars = sum(len(p.get_text()) for p in list(pdf)[: min(8, pdf.page_count)]) / min(
        8, pdf.page_count
    )
    return landscape > pdf.page_count / 2 or avg_chars < 900


def load_pdf_deck(
    path: str | Path,
    title: str | None = None,
    assets_dir: str | Path | None = None,
    dpi: int = 80,
    max_pages: int | None = None,
) -> Document:
    import pymupdf

    path = Path(path)
    doc_id = doc_id_for(str(path))
    pdf = pymupdf.open(str(path))
    segments: list[Segment] = []
    assets: dict[str, str] = {}
    if assets_dir:
        Path(assets_dir).mkdir(parents=True, exist_ok=True)
    for pno, page in enumerate(pdf, start=1):
        if max_pages and pno > max_pages:
            break
        txt = re.sub(r"[ \t]+", " ", page.get_text()).strip()
        txt = re.sub(r"\n{2,}", "\n", txt)
        asset_ids: list[str] = []
        if assets_dir:
            aid = f"{doc_id}_p{pno:03d}"
            out = Path(assets_dir) / f"{aid}.png"
            if not out.exists():
                page.get_pixmap(dpi=dpi).save(str(out))
            assets[aid] = str(out)
            asset_ids.append(aid)
        if txt:
            segments.append(
                Segment(
                    seg_id=f"s{len(segments):04d}",
                    text=txt,
                    kind=SegmentKind.SLIDE_TEXT,
                    assets=asset_ids,
                    span=SourceSpan(modality=Modality.SLIDES, doc_id=doc_id, slide_index=pno),
                )
            )
    meta_title = (pdf.metadata or {}).get("title") or ""
    return Document(
        doc_id=doc_id,
        title=title or meta_title or path.stem,
        modality=Modality.SLIDES,
        source=str(path),
        segments=segments,
        assets=assets,
        meta={"slides": pdf.page_count},
    )
