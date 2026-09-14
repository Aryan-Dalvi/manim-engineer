"""Presenter deck: one PPTX slide per beat with the section clip embedded (python-pptx)."""

from __future__ import annotations

from pathlib import Path

from manim_engineer.core import ArtifactBundle


def build_pptx(
    bundle: ArtifactBundle, bundle_dir: Path, out_path: Path, title: str = ""
) -> Path | None:
    try:
        from pptx import Presentation
        from pptx.util import Emu, Pt
    except ImportError:  # pragma: no cover
        return None
    prs = Presentation()
    prs.slide_width = Emu(12192000)  # 16:9
    prs.slide_height = Emu(6858000)
    blank = prs.slide_layouts[6]
    w, h = prs.slide_width, prs.slide_height
    for sec in bundle.sections:
        mp4 = bundle_dir / sec.mp4
        png = bundle_dir / sec.png if sec.png else None
        if not mp4.exists():
            continue
        slide = prs.slides.add_slide(blank)
        margin = int(w * 0.04)
        top = int(h * 0.12)
        vid_w = w - 2 * margin
        vid_h = int(vid_w * bundle.height / max(bundle.width, 1))
        if vid_h > h - top - margin:
            vid_h = h - top - margin
            vid_w = int(vid_h * bundle.width / max(bundle.height, 1))
        left = int((w - vid_w) / 2)
        slide.shapes.add_movie(
            str(mp4),
            left,
            top,
            vid_w,
            vid_h,
            poster_frame_image=str(png) if png and png.exists() else None,
            mime_type="video/mp4",
        )
        tb = slide.shapes.add_textbox(margin, int(h * 0.02), w - 2 * margin, int(h * 0.09))
        p = tb.text_frame.paragraphs[0]
        p.text = f"{title or bundle.concept_id}: {sec.name}"
        p.font.size = Pt(20)
        if sec.narration:
            slide.notes_slide.notes_text_frame.text = sec.narration
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    return out_path
